#!/usr/bin/env python3
"""
TARS Lab Academy — Day 01 deck built ON the green futuristic template.
Clones the template's branded slides (keeping the corner decorations, gradient
background and embedded Montserrat fonts), rebrands to TARS, and injects
expanded ~1-hour content (each datatype / operator / loop on its own slide).

Run:  python3 slides/build_pptx_v2.py
Out:  slides/TARS_Day01_Intro_Python_Core_v2.pptx
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

TEMPLATE = "Green Modern Futuristic Artificial Intelligence Presentation.pptx"
OUT = "slides/TARS_Day01_Intro_Python_Core_v2.pptx"

# ---------- palette (sampled from the template) ----------
INK    = RGBColor(0xFF, 0xFF, 0xFF)
MUTED  = RGBColor(0xCF, 0xDE, 0xC9)   # soft green-white
ACCENT = RGBColor(0x80, 0xC7, 0x72)   # signature light green
ACC2   = RGBColor(0xC2, 0xB8, 0x33)   # template yellow
TEAL   = RGBColor(0x54, 0xB0, 0x87)
CARD   = RGBColor(0x10, 0x2A, 0x12)   # panel green
BORDER = RGBColor(0x2C, 0x5A, 0x2E)
CODEBG = RGBColor(0x05, 0x12, 0x07)
DIM    = RGBColor(0x9C, 0xB3, 0x96)

FONT_B = "Montserrat Bold"
FONT_S = "Montserrat Semi-Bold"
MONO   = "Consolas"

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

prs = Presentation(TEMPLATE)
SW, SH = prs.slide_width, prs.slide_height
N_ORIG = len(prs.slides._sldIdLst)
SRC_TITLE   = prs.slides[0]   # "Artificial Intelligence PRESENTATION"
SRC_CONTENT = prs.slides[1]   # "About Us" — clean title+body chrome


# ============================================================
# slide duplication (preserves decorations, images, background)
# ============================================================
def dup_slide(src):
    layout = src.slide_layout
    dest = prs.slides.add_slide(layout)
    # strip placeholders that add_slide created
    spTree = dest.shapes._spTree
    for sp in list(spTree):
        if sp.tag.split('}')[-1] in ('sp', 'pic', 'grpSp', 'graphicFrame', 'cxnSp'):
            spTree.remove(sp)
    # copy gradient background
    src_bg = src._element.find(qn('p:cSld')).find(qn('p:bg'))
    if src_bg is not None:
        dest._element.find(qn('p:cSld')).insert(0, copy.deepcopy(src_bg))
    # remap relationships (images / svg)
    rid_map = {}
    for rId, rel in src.part.rels.items():
        rt = rel.reltype
        if rt.endswith('slideLayout') or rt.endswith('notesSlide'):
            continue
        if rel.is_external:
            rid_map[rId] = dest.part.relate_to(rel.target_ref, rt, is_external=True)
        else:
            rid_map[rId] = dest.part.relate_to(rel.target_part, rt)
    # copy shapes with rId remap
    for shp in src.shapes:
        el = copy.deepcopy(shp._element)
        for node in el.iter():
            for a in ('embed', 'link', 'id'):
                k = qn('r:' + a)
                v = node.get(k)
                if v in rid_map:
                    node.set(k, rid_map[v])
        spTree.append(el)
    return dest


def get_shape(slide, sid):
    for sp in slide.shapes:
        if sp.shape_id == sid:
            return sp
    return None


def remove_ids(slide, ids):
    for sp in list(slide.shapes):
        if sp.shape_id in ids:
            sp._element.getparent().remove(sp._element)


def set_box_text(shape, text, color=None, size=None, font=None):
    """Replace a textbox's text, keeping its first run's formatting."""
    tf = shape.text_frame
    p = tf.paragraphs[0]
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        run = runs[0]
    else:
        run = p.add_run(); run.text = text
    if color is not None: run.font.color.rgb = color
    if size is not None:  run.font.size = Pt(size)
    if font is not None:  run.font.name = font


# ---------- content drawing on cloned slides ----------
def _font(run, size, color, font=FONT_S, bold=False):
    run.font.size = Pt(size); run.font.color.rgb = color
    run.font.name = font; run.font.bold = bold


def tbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    return tb, tf


def para(tf, parts, size, color=INK, font=FONT_S, align=PP_ALIGN.LEFT,
         after=6, before=0, line=None, first=False):
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(after); p.space_before = Pt(before)
    if line is not None: p.line_spacing = line
    if isinstance(parts, str):
        parts = [(parts, color, font)]
    for tup in parts:
        t, c = tup[0], tup[1]
        f = tup[2] if len(tup) > 2 else font
        r = p.add_run(); r.text = t; _font(r, size, c, f)
    return p


def rrect(s, x, y, w, h, fill=None, line=None, lw=1.0, rad=True):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    return shp


# layout constants for 20 x 11.25 canvas
LX = Inches(1.72)            # left content margin
CW = Inches(16.9)            # content width
TITLE_Y = Inches(1.78)
BODY_Y = Inches(3.5)


def new_content(page, kicker, title_parts, title_size=46):
    s = dup_slide(SRC_CONTENT)
    remove_ids(s, {15, 16, 17, 18})          # drop template's About/Us + lorem
    b = get_shape(s, 13)                      # branding
    if b: set_box_text(b, "TARS LAB ACADEMY", color=ACCENT)
    pg = get_shape(s, 14)                     # page number
    if pg: set_box_text(pg, f"DAY 01  ·  {page:02d}")
    # kicker
    tb, tf = tbox(s, LX, Inches(1.18), CW, Inches(0.4))
    para(tf, [("—  " + kicker.upper(), ACCENT, MONO)], 12, first=True)
    # title
    tb, tf = tbox(s, LX, TITLE_Y, CW, Inches(1.5))
    para(tf, title_parts, title_size, font=FONT_B, line=1.02, first=True)
    return s


def bullets(s, x, y, w, items, size=17, gap=11):
    tb, tf = tbox(s, x, y, w, Inches(6))
    for i, parts in enumerate(items):
        para(tf, [("▹  ", ACCENT, FONT_B)] + parts, size, after=gap, line=1.12,
             first=(i == 0))


def codebox(s, x, y, w, h, lines, size=14, title=None):
    rrect(s, x, y, w, h, fill=CODEBG, line=BORDER, lw=1.25)
    if title:
        tb, tf = tbox(s, x + Inches(0.3), y + Inches(0.14), w - Inches(0.6), Inches(0.3))
        para(tf, [(title, DIM, MONO)], 11, first=True)
        ty = y + Inches(0.5)
    else:
        ty = y + Inches(0.24)
    tb, tf = tbox(s, x + Inches(0.32), ty, w - Inches(0.64), h - Inches(0.5))
    for i, segs in enumerate(lines):
        para(tf, [(t, c, MONO) for (t, c) in segs], size, after=3, line=1.12,
             first=(i == 0))


def card(s, x, y, w, h, head, body, icon=None, accent=ACCENT):
    rrect(s, x, y, w, h, fill=CARD, line=BORDER, lw=1.0)
    tb, tf = tbox(s, x + Inches(0.26), y + Inches(0.2), w - Inches(0.5), h - Inches(0.36))
    if icon:
        para(tf, [(icon, INK, FONT_S)], 20, first=True, after=4)
        para(tf, [(head, accent, FONT_B)], 15, after=4)
    else:
        para(tf, [(head, accent, FONT_B)], 15, first=True, after=4)
    para(tf, [(body, MUTED, FONT_S)], 11.5, line=1.18)


def table(s, x, y, w, rows, col_w, head_size=13, body_size=12.5, rh=Inches(0.46)):
    """rows[0] = header. col_w = list of fractions summing ~1."""
    cx = x
    cur_y = y
    for ri, row in enumerate(rows):
        cx = x
        is_head = (ri == 0)
        if is_head:
            rrect(s, x, cur_y, w, rh, fill=CARD, line=None)
        for ci, cell in enumerate(row):
            cwd = Emu(int(int(w) * col_w[ci]))
            tb, tf = tbox(s, cx + Inches(0.15), cur_y, cwd - Inches(0.2), rh,
                          anchor=MSO_ANCHOR.MIDDLE)
            if is_head:
                para(tf, [(cell, ACCENT, FONT_B)], head_size, first=True)
            else:
                c0 = INK if ci == 0 else MUTED
                f0 = MONO if ci == 0 else FONT_S
                para(tf, [(cell, c0, f0)], body_size, first=True, line=1.05)
            cx = cx + cwd
        # row divider
        ln = rrect(s, x, cur_y + rh - Pt(0.75), w, Pt(0.75), fill=BORDER, rad=False)
        cur_y = cur_y + rh
    return cur_y


def note(s, text, y=Inches(9.55)):
    tb, tf = tbox(s, LX, y, CW, Inches(0.9))
    para(tf, [("◆  ", ACC2, FONT_B), (text, MUTED, FONT_S)], 14, first=True, line=1.2)


# ============================================================
#  TITLE SLIDE
# ============================================================
s = dup_slide(SRC_TITLE)
# rewrite the three big title boxes + branding + page (ids from template slide1)
set_box_text(get_shape(s, 9),  "Generative AI",  color=INK)
set_box_text(get_shape(s, 10), "Engineer Bootcamp", color=ACCENT)
set_box_text(get_shape(s, 11), "DAY 01  ·  PYTHON CORE & ENVIRONMENTS")
set_box_text(get_shape(s, 12), "TARS LAB ACADEMY", color=ACCENT)
pg = get_shape(s, 16)
if pg: set_box_text(pg, "45-DAY PROGRAM")

# ============================================================
#  INTRO
# ============================================================
s = new_content(1, "Welcome", [("Welcome to ", INK), ("TARS Lab Academy", ACCENT)], 40)
tb, tf = tbox(s, LX, BODY_Y - Inches(0.6), CW, Inches(1.0))
para(tf, [("We don't teach you to watch — we teach you to ", MUTED),
          ("ship", ACCENT, FONT_B),
          (". Every day pairs one focused concept with a real build, so by Day 45 "
           "you have a portfolio, not just notes.", MUTED)], 16, first=True, line=1.3)
cw3, gap = Inches(5.4), Inches(0.35)
cy, ch = Inches(5.4), Inches(2.7)
card(s, LX, cy, cw3, ch, "Outcome-first", "Each day ends with something running on your machine.", "🎯")
card(s, LX+cw3+gap, cy, cw3, ch, "Foundations → Frontier", "Python today; RAG, agents, fine-tuning & LLMOps later.", "🧱")
card(s, LX+2*(cw3+gap), cy, cw3, ch, "Industry tools", "LangChain, LangGraph, AutoGen, MCP, Docker, K8s.", "🛠️")

# --- the 45-day map ---
s = new_content(2, "The 45-day map", [("12 phases, one journey", INK)], 42)
phases = [
    ("P1", "Foundations (Python)", "D1–3"), ("P2", "How LLMs Work", "D4–7"),
    ("P3", "Building LLM Apps", "D8–11"), ("P4", "Advanced RAG", "D12–16"),
    ("P5", "AI Agents", "D17–21"), ("P6", "Multi-Agent (AutoGen)", "D22–25"),
    ("P7", "MCP & Coding Tools", "D26–29"), ("P8", "Fine-Tuning", "D30–32"),
    ("P9", "Evaluation", "D33–35"), ("P10", "Security & Guardrails", "D36–38"),
    ("P11", "LLMOps & Production", "D39–42"), ("P12", "Capstone & Career", "D43–45"),
]
colw, rowh = Inches(8.2), Inches(0.92)
x0, y0 = LX, Inches(3.3)
for idx, (pn, pt, pd) in enumerate(phases):
    col, row = idx % 2, idx // 2
    x = x0 + col * (colw + Inches(0.5)); y = y0 + row * rowh
    rrect(s, x, y + rowh - Pt(2), colw, Pt(1), fill=BORDER, rad=False)
    tb, tf = tbox(s, x, y, colw, rowh - Inches(0.1), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(pn + "   ", ACCENT, MONO), (pt, INK, FONT_S)], 15, first=True)
    tb, tf = tbox(s, x, y, colw, rowh - Inches(0.1), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(pd, ACC2, MONO)], 13, align=PP_ALIGN.RIGHT, first=True)

# --- how every day works ---
s = new_content(3, "The rhythm", [("How every day works", INK)], 42)
cw3, gap, cy, ch = Inches(5.4), Inches(0.35), Inches(3.6), Inches(3.0)
card(s, LX, cy, cw3, ch, "1 · Learn", "A tight concept block — theory + intuition, no fluff.", "📖")
card(s, LX+cw3+gap, cy, cw3, ch, "2 · Code along", "A guided Jupyter notebook you run line by line with us.", "⌨️")
card(s, LX+2*(cw3+gap), cy, cw3, ch, "3 · Build", "A daily challenge you finish on your own.", "🔨")
note(s, "Today's notebook:  Day01_Python_Core_and_Environments.ipynb", y=Inches(7.1))

# ============================================================
#  DAY 1 DIVIDER
# ============================================================
s = new_content(4, "Phase 1 · Foundations", [("Day 01", ACCENT)], 96)
tb, tf = tbox(s, LX, Inches(5.4), CW, Inches(1.0))
para(tf, [("Python Core & Environments", INK, FONT_B)], 38, first=True)
tb, tf = tbox(s, LX, Inches(6.5), Inches(13), Inches(1.2))
para(tf, [("The language every GenAI tool speaks — and the workspace you'll live "
           "in for the next 45 days.", MUTED)], 17, first=True, line=1.3)

# --- agenda ---
s = new_content(5, "Today's agenda", [("What we cover in this hour", INK)], 38)
ag = [("⚙️", "Environments", "VS Code · Conda · UV · virtual environments"),
      ("🔤", "Variables & types", "int, float, complex, str, bool, None + casting"),
      ("➗", "Operators", "arithmetic, comparison, logical, assignment, membership"),
      ("🔁", "Control flow", "if / elif / else · for · while · break / continue"),
      ("📦", "Data structures", "list, comprehensions, tuple, set, dict"),
      ("🔨", "Build", "data-manipulation exercises in the notebook")]
cw2, gap2, ch2 = Inches(8.3), Inches(0.3), Inches(1.55)
for i, (ic, t, b) in enumerate(ag):
    col, row = i % 2, i // 2
    card(s, LX + col*(cw2+gap2), Inches(3.35) + row*(ch2+Inches(0.22)), cw2, ch2, t, b, ic)

# ============================================================
#  ENVIRONMENTS
# ============================================================
s = new_content(6, "Foundation", [("Why Python for GenAI?", INK)], 44)
bullets(s, LX, Inches(3.5), CW, [
    [("It's the lingua franca", ACCENT, FONT_B),
     (" — PyTorch, Transformers, LangChain — every SDK is Python-first.", INK)],
    [("Fast to prototype", ACCENT, FONT_B),
     (" — readable syntax takes you from idea to running code in minutes.", INK)],
    [("Massive ecosystem", ACCENT, FONT_B),
     (" — data, ML, web and deployment all live in one language.", INK)],
    [("The glue layer", ACCENT, FONT_B),
     (" — call APIs, orchestrate agents and serve models from one script.", INK)],
], size=19, gap=18)

s = new_content(7, "Setup", [("Your workspace: VS Code · Conda · UV", INK)], 34)
cw3, gap, cy, ch = Inches(5.4), Inches(0.35), Inches(3.6), Inches(3.2)
card(s, LX, cy, cw3, ch, "VS Code", "Editor + Python & Jupyter extensions. Where you write, run and debug everything.", "🧩")
card(s, LX+cw3+gap, cy, cw3, ch, "Conda", "Manages Python versions and heavyweight ML packages cleanly, in isolated environments.", "🐍")
card(s, LX+2*(cw3+gap), cy, cw3, ch, "UV", "Ultra-fast package & project manager — like pip, but resolves in seconds not minutes.", "⚡")
note(s, "Tip: Ctrl/Cmd+Shift+P → 'Python: Select Interpreter' to point VS Code at your env.", y=Inches(7.1))

s = new_content(8, "Isolation", [("Virtual environments — why & how", INK)], 36)
tb, tf = tbox(s, LX, Inches(3.0), CW, Inches(0.8))
para(tf, [("Each project gets its own sandbox of dependencies, so versions never "
           "collide. Non-negotiable in AI work.", MUTED)], 15, first=True, line=1.25)
codebox(s, LX, Inches(3.9), CW, Inches(4.6), [
    [("# Option A — Conda", DIM)],
    [("conda create ", INK), ("-n ", ACC2), ("genai ", TEAL), ("python=3.11 ", INK), ("-y", ACC2)],
    [("conda activate ", INK), ("genai", TEAL)],
    [("", INK)],
    [("# Option B — UV  (fast, modern)", DIM)],
    [("uv venv", INK), ("                          # creates .venv", DIM)],
    [("source .venv/bin/activate", INK), ("     # Windows: .venv\\Scripts\\activate", DIM)],
    [("uv pip install ", INK), ("numpy pandas streamlit", TEAL)],
    [("", INK)],
    [("# verify which Python you're in", DIM)],
    [("python ", INK), ("-c ", ACC2), ('"import sys; print(sys.executable)"', TEAL)],
], size=14.5, title="terminal")

# ============================================================
#  VARIABLES & DATA TYPES  (one per slide)
# ============================================================
s = new_content(9, "Language core", [("Variables & dynamic typing", INK)], 40)
tb, tf = tbox(s, LX, Inches(3.0), CW, Inches(0.9))
para(tf, [("No type declarations — a variable is just a ", MUTED),
          ("name bound to an object", ACCENT, FONT_B),
          (". The type lives on the value, and can change.", MUTED)], 15, first=True, line=1.25)
codebox(s, LX, Inches(4.0), CW, Inches(4.3), [
    [("name ", INK), ("= ", ACCENT), ('"TARS"', ACC2), ("        # str", DIM)],
    [("days ", INK), ("= ", ACCENT), ("45", TEAL), ("             # int", DIM)],
    [("days ", INK), ("= ", ACCENT), ('"forty-five"', ACC2), ("   # now a str — totally legal", DIM)],
    [("", INK)],
    [("x = y = z ", INK), ("= ", ACCENT), ("0", TEAL), ("        # chained assignment", DIM)],
    [("a, b ", INK), ("= ", ACCENT), ("1, 2", TEAL), ("          # multiple assignment", DIM)],
    [("a, b ", INK), ("= ", ACCENT), ("b, a", TEAL), ("          # swap, no temp var", DIM)],
    [("", INK)],
    [("print(", INK), ("type", ACCENT), ("(name), ", INK), ("type", ACCENT), ("(days))", INK),
     ("   # <class 'str'> ...", DIM)],
], size=14.5, title="python")

# numbers
s = new_content(10, "Data types · 1 of 4", [("Numbers — ", INK), ("int · float · complex", ACCENT)], 36)
codebox(s, LX, Inches(3.2), CW, Inches(3.2), [
    [("count   ", INK), ("= ", ACCENT), ("42", TEAL), ("            # int  (unlimited size)", DIM)],
    [("price   ", INK), ("= ", ACCENT), ("19.99", TEAL), ("         # float", DIM)],
    [("big     ", INK), ("= ", ACCENT), ("1_000_000", TEAL), ("     # underscores for readability", DIM)],
    [("z       ", INK), ("= ", ACCENT), ("2 + 3j", TEAL), ("        # complex", DIM)],
    [("ratio   ", INK), ("= ", ACCENT), ("int", ACCENT), ("(", INK), ("price", TEAL), (")", INK),
     ("     # -> 19  (truncates)", DIM)],
], size=14.5, title="python")
table(s, LX, Inches(6.7), CW, [
    ["Operation", "Example", "Result"],
    ["round / abs", "round(19.99, 1) · abs(-4)", "20.0 · 4"],
    ["int ↔ float", "float(3) · int(3.9)", "3.0 · 3"],
    ["divmod", "divmod(17, 5)", "(3, 2)  → quotient, remainder"],
], col_w=[0.28, 0.42, 0.30])

# strings
s = new_content(11, "Data types · 2 of 4", [("Strings — ", INK), ("text", ACCENT)], 38)
codebox(s, LX, Inches(3.2), CW, Inches(3.6), [
    [("course ", INK), ("= ", ACCENT), ('"GenAI Bootcamp"', ACC2)],
    [("course[", INK), ("0", TEAL), ("]      ", INK), ("# 'G'   indexing", DIM)],
    [("course[", INK), ("0:5", TEAL), ("]    ", INK), ("# 'GenAI'   slicing", DIM)],
    [("course[::-", INK), ("1", TEAL), ("]   ", INK), ("# reversed", DIM)],
    [("course.upper(), course.split(), course.replace(", INK), ('"AI"', ACC2), (", ", INK), ('"ML"', ACC2), (")", INK)],
    [('f"', ACC2), ("Day ", ACC2), ("{", INK), ("1", TEAL), ("}", INK), (" of ", ACC2), ("{", INK), ("45", TEAL), ("}", INK), ('"', ACC2),
     ("           # f-string -> 'Day 1 of 45'", DIM)],
    [('f"', ACC2), ("{", INK), ("price", TEAL), (":.2f", ACC2), ("}", INK), ('"', ACC2),
     ("              # format spec -> '19.99'", DIM)],
], size=14, title="python")
note(s, "Strings are immutable — methods return a NEW string, they never modify in place.", y=Inches(7.2))

# booleans & none
s = new_content(12, "Data types · 3 of 4", [("Booleans & ", INK), ("None", ACCENT)], 38)
codebox(s, LX, Inches(3.2), CW, Inches(2.9), [
    [("is_ready ", INK), ("= ", ACCENT), ("True", TEAL), ("          # bool: True / False", DIM)],
    [("result   ", INK), ("= ", ACCENT), ("None", TEAL), ("          # 'no value yet'", DIM)],
    [("bool(", ACCENT), ("0", TEAL), ("), bool(", ACCENT), ('""', ACC2), ("), bool(", ACCENT), ("[]", INK), (")",INK),
     ("   # all False (falsy)", DIM)],
    [("bool(", ACCENT), ("42", TEAL), ("), bool(", ACCENT), ('"hi"', ACC2), (")", INK),
     ("        # both True (truthy)", DIM)],
], size=14.5, title="python")
bullets(s, LX, Inches(6.5), CW, [
    [("Falsy values:", ACC2, FONT_B), (" 0, 0.0, '', [], {}, set(), None, False", MUTED, MONO)],
    [("Everything else is truthy", ACCENT, FONT_B), (" — used directly in if-statements.", INK)],
], size=15, gap=12)

# casting / type checks
s = new_content(13, "Data types · 4 of 4", [("Type conversion & checking", INK)], 36)
codebox(s, LX, Inches(3.2), CW, Inches(3.2), [
    [("int(", ACCENT), ('"10"', ACC2), ("), float(", ACCENT), ('"3.14"', ACC2), ("), str(", ACCENT), ("99", TEAL), (")", INK)],
    [("list(", ACCENT), ('"abc"', ACC2), ("), set([", ACCENT), ("1", TEAL), (",", INK), ("1", TEAL), (",", INK), ("2", TEAL), ("]), tuple([", ACCENT), ("1", TEAL), (",", INK), ("2", TEAL), ("])", INK)],
    [("", INK)],
    [("type", ACCENT), ("(x) ", INK), ("is ", ACCENT), ("int", ACCENT), ("              # exact check", DIM)],
    [("isinstance", ACCENT), ("(x, (", INK), ("int", ACCENT), (", ", INK), ("float", ACCENT), ("))", INK),
     ("    # preferred — allows subtypes", DIM)],
], size=14, title="python")
note(s, "Rule of thumb: use isinstance() for checks; bad casts raise ValueError — handle them (Day 2).", y=Inches(7.0))

# ============================================================
#  OPERATORS  (expanded)
# ============================================================
s = new_content(14, "Operators · 1 of 3", [("Arithmetic operators", INK)], 40)
table(s, LX, Inches(3.3), CW, [
    ["Operator", "Meaning", "Example → Result"],
    ["+  -  *", "add · subtract · multiply", "7 + 2 → 9"],
    ["/", "true division (float)", "7 / 2 → 3.5"],
    ["//", "floor division", "7 // 2 → 3"],
    ["%", "modulo (remainder)", "7 % 2 → 1"],
    ["**", "exponent (power)", "7 ** 2 → 49"],
], col_w=[0.22, 0.40, 0.38], rh=Inches(0.62))
note(s, "Precedence: **  →  * / // %  →  + -.  Use ( ) to be explicit and readable.", y=Inches(8.1))

s = new_content(15, "Operators · 2 of 3", [("Comparison & logical", INK)], 40)
table(s, LX, Inches(3.2), Inches(8.1), [
    ["Comparison", "Result"],
    ["==   !=", "equal · not equal"],
    ["<  >  <=  >=", "ordering"],
], col_w=[0.5, 0.5], rh=Inches(0.55))
table(s, Inches(10.3), Inches(3.2), Inches(8.3), [
    ["Logical", "Result"],
    ["and", "True if BOTH true"],
    ["or", "True if EITHER true"],
    ["not", "inverts the boolean"],
], col_w=[0.4, 0.6], rh=Inches(0.55))
codebox(s, LX, Inches(5.7), CW, Inches(2.4), [
    [("passed ", INK), ("= ", ACCENT), ("score >= ", INK), ("60 ", TEAL), ("and ", ACCENT), ("attendance > ", INK), ("0.8", TEAL)],
    [("18 ", TEAL), ("<= ", ACCENT), ("age ", INK), ("< ", ACCENT), ("65", TEAL),
     ("            # chained comparison — Pythonic", DIM)],
    [("result ", INK), ("= ", ACCENT), ("value ", INK), ("or ", ACCENT), ('"default"', ACC2),
     ("    # short-circuit fallback", DIM)],
], size=14, title="python")

s = new_content(16, "Operators · 3 of 3", [("Assignment, membership & identity", INK)], 32)
table(s, LX, Inches(3.2), CW, [
    ["Group", "Operators", "Example"],
    ["Assignment", "=  +=  -=  *=  /=  //=  **=", "total += 5"],
    ["Membership", "in   ·   not in", '"a" in "cat" → True'],
    ["Identity", "is   ·   is not", "x is None"],
], col_w=[0.24, 0.46, 0.30], rh=Inches(0.62))
note(s, "Use '==' to compare VALUES, 'is' only for identity (mainly 'is None'). They are not the same!", y=Inches(7.6))

# ============================================================
#  CONTROL FLOW  (expanded)
# ============================================================
s = new_content(17, "Control flow · 1 of 3", [("Conditionals — ", INK), ("if / elif / else", ACCENT)], 34)
tb, tf = tbox(s, LX, Inches(3.0), CW, Inches(0.6))
para(tf, [("Blocks are defined by ", MUTED), ("indentation", ACCENT, FONT_B),
          (" (4 spaces) — there are no braces.", MUTED)], 15, first=True)
codebox(s, LX, Inches(3.8), CW, Inches(4.4), [
    [("def ", ACCENT), ("phase_for", TEAL), ("(day):", INK)],
    [("    if ", ACCENT), ("day <= ", INK), ("3", TEAL), (":", INK)],
    [("        return ", ACCENT), ('"Foundations"', ACC2)],
    [("    elif ", ACCENT), ("day <= ", INK), ("7", TEAL), (":", INK)],
    [("        return ", ACCENT), ('"How LLMs Work"', ACC2)],
    [("    else", ACCENT), (":", INK)],
    [("        return ", ACCENT), ('"Advanced"', ACC2)],
    [("", INK)],
    [("status ", INK), ("= ", ACCENT), ('"pass" ', ACC2), ("if ", ACCENT), ("score >= ", INK), ("60 ", TEAL), ("else ", ACCENT), ('"fail"', ACC2),
     ("   # ternary", DIM)],
], size=14.5, title="python")

s = new_content(18, "Control flow · 2 of 3", [("Loops — ", INK), ("for & range", ACCENT)], 38)
codebox(s, LX, Inches(3.2), CW, Inches(4.2), [
    [("for ", ACCENT), ("day ", INK), ("in ", ACCENT), ("range", ACCENT), ("(", INK), ("1", TEAL), (", ", INK), ("46", TEAL), ("):", INK),
     ("       # 1..45", DIM)],
    [("    print(day, phase_for(day))", INK)],
    [("", INK)],
    [("for ", ACCENT), ("i, topic ", INK), ("in ", ACCENT), ("enumerate", ACCENT), ("(topics):", INK),
     ("   # index + value", DIM)],
    [("for ", ACCENT), ("k, v ", INK), ("in ", ACCENT), ("student.items():", INK),
     ("       # dict pairs", DIM)],
    [("for ", ACCENT), ("a, b ", INK), ("in ", ACCENT), ("zip", ACCENT), ("(names, scores):", INK),
     ("    # parallel iterate", DIM)],
], size=14.5, title="python")
note(s, "range(start, stop, step) is lazy — it generates numbers on demand, never a full list in memory.", y=Inches(7.7))

s = new_content(19, "Control flow · 3 of 3", [("Loops — ", INK), ("while · break · continue", ACCENT)], 34)
codebox(s, LX, Inches(3.2), CW, Inches(4.4), [
    [("day ", INK), ("= ", ACCENT), ("0", TEAL)],
    [("while ", ACCENT), ("True", TEAL), (":", INK)],
    [("    day += ", INK), ("1", TEAL)],
    [("    if ", ACCENT), ("day % ", INK), ("2 ", TEAL), ("== ", ACCENT), ("0", TEAL), (":", INK)],
    [("        continue", ACCENT), ("        # skip the rest, next loop", DIM)],
    [("    if ", ACCENT), ("day > ", INK), ("7", TEAL), (":", INK)],
    [("        break", ACCENT), ("           # exit the loop entirely", DIM)],
    [("    print(", INK), ('"odd study day"', ACC2), (", day)", INK)],
], size=14.5, title="python")
note(s, "for/while can have an else: that runs only if the loop finished WITHOUT a break.", y=Inches(7.9))

# ============================================================
#  DATA STRUCTURES  (one per slide)
# ============================================================
s = new_content(20, "Data structures · list", [("Lists — ", INK), ("ordered, mutable", ACCENT)], 36)
codebox(s, LX, Inches(3.2), CW, Inches(3.8), [
    [("topics ", INK), ("= ", ACCENT), ("[", INK), ('"setup"', ACC2), (", ", INK), ('"vars"', ACC2), (", ", INK), ('"loops"', ACC2), ("]", INK)],
    [("topics.append(", INK), ('"dicts"', ACC2), (")        ", INK), ("# add to end", DIM)],
    [("topics.insert(", INK), ("0", TEAL), (", ", INK), ('"intro"', ACC2), (")     ", INK), ("# add at index", DIM)],
    [("topics[", INK), ("0", TEAL), ("] ", INK), ("= ", ACCENT), ('"env"', ACC2), ("           ", INK), ("# update in place", DIM)],
    [("topics.pop(), topics.remove(", INK), ('"loops"', ACC2), (")", INK)],
    [("topics[", INK), ("1", TEAL), (":", INK), ("3", TEAL), ("], len(topics), sorted(topics)", INK)],
], size=14, title="python")
note(s, "Lists hold mixed types and grow/shrink freely — the default container you reach for.", y=Inches(7.5))

s = new_content(21, "Data structures · comprehensions", [("List comprehensions", INK)], 38)
tb, tf = tbox(s, LX, Inches(3.0), CW, Inches(0.6))
para(tf, [("Build a new list in one line:  ", MUTED),
          ("[ expr for item in iterable if condition ]", ACCENT, MONO)], 15, first=True)
codebox(s, LX, Inches(3.9), CW, Inches(3.6), [
    [("squares ", INK), ("= ", ACCENT), ("[n", INK), ("**", ACCENT), ("2 ", TEAL), ("for ", ACCENT), ("n ", INK), ("in ", ACCENT), ("range", ACCENT), ("(", INK), ("10", TEAL), (")]", INK)],
    [("evens   ", INK), ("= ", ACCENT), ("[n ", INK), ("for ", ACCENT), ("n ", INK), ("in ", ACCENT), ("nums ", INK), ("if ", ACCENT), ("n % ", INK), ("2 ", TEAL), ("== ", ACCENT), ("0", TEAL), ("]", INK)],
    [("labels  ", INK), ("= ", ACCENT), ("[t.upper() ", INK), ("for ", ACCENT), ("t ", INK), ("in ", ACCENT), ("topics]", INK)],
    [("", INK)],
    [("# also: dict & set comprehensions", DIM)],
    [("{w: ", INK), ("len", ACCENT), ("(w) ", INK), ("for ", ACCENT), ("w ", INK), ("in ", ACCENT), ("topics}", INK), ("   # dict", DIM)],
], size=14, title="python")
note(s, "Comprehensions are faster and clearer than building a list with an empty list + append loop.", y=Inches(7.8))

s = new_content(22, "Data structures · tuple", [("Tuples — ", INK), ("ordered, immutable", ACCENT)], 36)
codebox(s, LX, Inches(3.2), CW, Inches(3.4), [
    [("point ", INK), ("= ", ACCENT), ("(", INK), ("10", TEAL), (", ", INK), ("20", TEAL), (")", INK), ("          # fixed pair", DIM)],
    [("x, y ", INK), ("= ", ACCENT), ("point", INK), ("            # unpacking", DIM)],
    [("rgb ", INK), ("= ", ACCENT), ("255", TEAL), (", ", INK), ("128", TEAL), (", ", INK), ("0", TEAL), ("          # parens optional", DIM)],
    [("point[", INK), ("0", TEAL), ("] ", INK), ("= ", ACCENT), ("5", TEAL),
     ("            # ❌ TypeError — immutable", DIM)],
], size=14.5, title="python")
bullets(s, LX, Inches(6.7), CW, [
    [("Use a tuple when data ", INK), ("shouldn't change", ACCENT, FONT_B),
     (" — coordinates, RGB, DB rows, dict keys.", INK)],
], size=15, gap=8)

s = new_content(23, "Data structures · set", [("Sets — ", INK), ("unique, unordered", ACCENT)], 36)
codebox(s, LX, Inches(3.2), CW, Inches(3.2), [
    [("tools ", INK), ("= ", ACCENT), ("{", INK), ('"uv"', ACC2), (", ", INK), ('"conda"', ACC2), (", ", INK), ('"uv"', ACC2), ("}", INK), ("   # -> {'uv', 'conda'}", DIM)],
    [('"uv" ', ACC2), ("in ", ACCENT), ("tools", INK), ("               # fast membership test", DIM)],
    [("set(", ACCENT), ("[", INK), ("1", TEAL), (",", INK), ("1", TEAL), (",", INK), ("2", TEAL), (",", INK), ("3", TEAL), ("])", INK), ("              # de-duplicate a list", DIM)],
    [("a | b   a & b   a - b", INK), ("       # union, intersect, difference", DIM)],
], size=14.5, title="python")
note(s, "Reach for a set to remove duplicates or test membership across thousands of items instantly.", y=Inches(7.0))

s = new_content(24, "Data structures · dict", [("Dictionaries — ", INK), ("key → value", ACCENT)], 34)
codebox(s, LX, Inches(3.2), CW, Inches(4.0), [
    [("student ", INK), ("= ", ACCENT), ("{", INK), ('"name"', ACC2), (": ", INK), ('"Ada"', ACC2), (", ", INK), ('"scores"', ACC2), (": [", INK), ("82", TEAL), (",", INK), ("91", TEAL), ("]}", INK)],
    [("student[", INK), ('"name"', ACC2), ("]                ", INK), ("# 'Ada'", DIM)],
    [("student.get(", INK), ('"email"', ACC2), (", ", INK), ('"n/a"', ACC2), (")     ", INK), ("# safe lookup", DIM)],
    [("student[", INK), ('"city"', ACC2), ("] ", INK), ("= ", ACCENT), ('"Lagos"', ACC2), ("       ", INK), ("# add / update", DIM)],
    [("for ", ACCENT), ("k, v ", INK), ("in ", ACCENT), ("student.items():", INK)],
    [("    print(k, v)", INK)],
], size=14, title="python")
note(s, "Dicts are the backbone of Python — JSON, API responses and config all map straight to dicts.", y=Inches(7.6))

# ============================================================
#  BUILD + RECAP + CLOSE
# ============================================================
s = new_content(25, "🔨 Today's build", [("Data-manipulation exercises", INK)], 38)
tb, tf = tbox(s, LX, Inches(3.0), CW, Inches(0.7))
para(tf, [("Open the notebook and work through it with us — then finish the challenge "
           "cells on your own.", MUTED)], 16, first=True, line=1.3)
bullets(s, LX, Inches(3.9), CW, [
    [("Clean a list of messy student records (strip, capitalize).", INK)],
    [("Add a computed average to each record.", INK)],
    [("De-duplicate cities with a ", INK), ("set", ACCENT, FONT_B), (".", INK)],
    [("Find top performers with a ", INK), ("comprehension", ACCENT, FONT_B), (".", INK)],
    [("Group students by city into a ", INK), ("dict", ACCENT, FONT_B), (".", INK)],
    [("Stretch: ", ACC2, FONT_B), ("count days per phase from the syllabus.", INK)],
], size=17, gap=12)
note(s, "File:  notebooks/Day01_Python_Core_and_Environments.ipynb", y=Inches(8.7))

s = new_content(26, "Recap", [("What you can do now", INK)], 40)
bullets(s, LX, Inches(3.4), CW, [
    [("Create isolated environments with ", INK), ("Conda / UV", ACCENT, FONT_B), (" + VS Code.", INK)],
    [("Use every core ", INK), ("datatype", ACCENT, FONT_B), (" and ", INK), ("operator", ACCENT, FONT_B), (" with confidence.", INK)],
    [("Direct program flow with ", INK), ("conditionals and loops", ACCENT, FONT_B), (".", INK)],
    [("Pick the right ", INK), ("data structure", ACCENT, FONT_B), (" and write comprehensions.", INK)],
], size=18, gap=16)
note(s, "Homework: finish the notebook challenges and push it to your GitHub repo.", y=Inches(8.0))

# closing — clone title for impact
s = dup_slide(SRC_TITLE)
set_box_text(get_shape(s, 9),  "Let's", color=INK)
set_box_text(get_shape(s, 10), "build.", color=ACCENT)
set_box_text(get_shape(s, 11), "NEXT — DAY 02 · PYTHON OOP & ADVANCED")
set_box_text(get_shape(s, 12), "TARS LAB ACADEMY", color=ACCENT)
pg = get_shape(s, 16)
if pg: set_box_text(pg, "SEE YOU TOMORROW")

# ============================================================
#  delete the 15 original template slides
# ============================================================
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst)[:N_ORIG]:
    rId = sldId.get(qn('r:id'))
    prs.part.drop_rel(rId)
    sldIdLst.remove(sldId)

prs.save(OUT)
print(f"✅ saved {OUT} — {len(prs.slides._sldIdLst)} slides")
