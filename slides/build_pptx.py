#!/usr/bin/env python3
"""
TARS Lab Academy — PPTX generator (Day 01 · Intro & Python Core)
Builds an editable, on-brand PowerPoint deck matching slides/day01-intro.html.

Run:  python3 slides/build_pptx.py
Out:  slides/TARS_Day01_Intro_Python_Core.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- TARS brand palette ----------
BG      = RGBColor(0x0A, 0x0E, 0x14)
BG2     = RGBColor(0x0F, 0x16, 0x22)
PANEL   = RGBColor(0x12, 0x1A, 0x28)
INK     = RGBColor(0xE8, 0xED, 0xF4)
MUTED   = RGBColor(0x8A, 0x97, 0xA8)
LINE    = RGBColor(0x1D, 0x27, 0x35)
ACCENT  = RGBColor(0x00, 0xD4, 0xC8)   # TARS cyan
AMBER   = RGBColor(0xFF, 0xB0, 0x2E)
INDIGO  = RGBColor(0x6C, 0x8C, 0xFF)
CODEBG  = RGBColor(0x06, 0x0A, 0x10)

FONT    = "Segoe UI"
MONO    = "Consolas"

# 16:9
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


# ---------- helpers ----------
def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    # send to back
    sp = bg._element
    sp.getparent().remove(sp)
    s.shapes._spTree.insert(2, sp)
    return s


def _set_font(run, size, color, bold=False, font=FONT, italic=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def textbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    return tb, tf


def para(tf, parts, size, color=INK, bold=False, space_after=6, space_before=0,
         align=PP_ALIGN.LEFT, font=FONT, line=None, first=False):
    """parts: str  OR  list of (text, color, bold[, font]) tuples."""
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if line is not None:
        p.line_spacing = line
    if isinstance(parts, str):
        parts = [(parts, color, bold, font)]
    for tup in parts:
        text, c, b = tup[0], tup[1], tup[2]
        f = tup[3] if len(tup) > 3 else font
        r = p.add_run(); r.text = text
        _set_font(r, size, c, b, f)
    return p


def rect(s, x, y, w, h, fill=None, line_color=None, line_w=1.0, radius=True):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def kicker(s, text, x=Inches(0.9), y=Inches(0.95)):
    tb, tf = textbox(s, x, y, Inches(11), Inches(0.4))
    para(tf, [("—  " + text.upper(), ACCENT, True, MONO)], 13, first=True)
    return tb


def brand(s):
    # logo mark: 3x3 grid, diagonal cyan
    gx, gy, c, gap = Inches(0.9), Inches(0.42), Inches(0.085), Inches(0.03)
    for r in range(3):
        for col in range(3):
            on = (r == col)
            sq = rect(s, gx + col*(c+gap), gy + r*(c+gap), c, c,
                      fill=ACCENT if on else LINE, radius=True)
    tb, tf = textbox(s, Inches(1.55), Inches(0.4), Inches(6), Inches(0.6))
    para(tf, [("TARS", ACCENT, True), (" LAB ACADEMY", INK, True)], 12, first=True,
         space_after=0)
    para(tf, [("GENERATIVE AI ENGINEER BOOTCAMP", MUTED, False, MONO)], 8.5,
         space_after=0)


def footer(s, n, total=16):
    tb, tf = textbox(s, Inches(0.9), Inches(7.02), Inches(7), Inches(0.35))
    para(tf, [("TARS Lab Academy  ·  GenAI Engineer Bootcamp", MUTED, False)], 9,
         first=True)
    tb2, tf2 = textbox(s, Inches(11.2), Inches(7.02), Inches(1.23), Inches(0.35))
    para(tf2, [(f"{n:02d}", ACCENT, True, MONO), (f" / {total:02d}", MUTED, False, MONO)],
         9, align=PP_ALIGN.RIGHT, first=True)
    # top accent bar (progress)
    w = int(SW * n / total)
    rect(s, 0, 0, Emu(w), Pt(3.2), fill=ACCENT, radius=False)


def chrome(s, n):
    brand(s); footer(s, n)


def card(s, x, y, w, h, icon, title, body):
    c = rect(s, x, y, w, h, fill=BG2, line_color=LINE, line_w=1.0)
    tb, tf = textbox(s, x + Inches(0.22), y + Inches(0.18),
                     w - Inches(0.44), h - Inches(0.36))
    para(tf, [(icon, INK, False)], 22, first=True, space_after=4)
    para(tf, [(title, INK, True)], 15, space_after=4)
    para(tf, [(body, MUTED, False)], 11.5, line=1.15)
    return c


def bullets(s, x, y, w, items, size=18, gap=10):
    tb, tf = textbox(s, x, y, w, Inches(4))
    for i, parts in enumerate(items):
        run_parts = [("▹  ", ACCENT, True)]
        run_parts += parts
        para(tf, run_parts, size, space_after=gap, line=1.1, first=(i == 0))


def codebox(s, x, y, w, h, lines):
    """lines: list of list-of-(text,color) segments."""
    rect(s, x, y, w, h, fill=CODEBG, line_color=LINE, line_w=1.0)
    tb, tf = textbox(s, x + Inches(0.28), y + Inches(0.2),
                     w - Inches(0.56), h - Inches(0.4))
    for i, segs in enumerate(lines):
        para(tf, [(t, c, False, MONO) for (t, c) in segs], 13.5,
             space_after=3, line=1.1, first=(i == 0))


def title_block(s, x, y, w, parts_lines, size=44, line=1.02):
    tb, tf = textbox(s, x, y, w, Inches(2.2))
    for i, parts in enumerate(parts_lines):
        para(tf, parts, size, bold=True, line=line, space_after=0, first=(i == 0))
    return tb


# ============================================================
# SLIDE 1 — TITLE
# ============================================================
s = slide()
# subtle accent glow panels
rect(s, Inches(9.2), Inches(-1.2), Inches(5), Inches(4), fill=BG2, radius=True)
brand(s)
kicker(s, "45-Day Bootcamp · Day 01", y=Inches(1.7))
title_block(s, Inches(0.9), Inches(2.2), Inches(11.5), [
    [("Become a ", INK, True), ("Generative AI", ACCENT, True)],
    [("Engineer.", INK, True)],
], size=54)
tb, tf = textbox(s, Inches(0.9), Inches(4.5), Inches(9.5), Inches(1.2))
para(tf, [("From Python fundamentals to production-grade RAG, agents, fine-tuning, "
           "evaluation, security and LLMOps — built one day at a time.", MUTED, False)],
     18, first=True, line=1.35)
# pills
px = Inches(0.9)
for label in ["12 Phases", "45 Days", "Build every single day"]:
    pw = Inches(0.34 + 0.105 * len(label))
    rect(s, px, Inches(5.85), pw, Inches(0.5), fill=BG2, line_color=ACCENT, line_w=1.0)
    tb, tf = textbox(s, px, Inches(5.85), pw, Inches(0.5), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(label, ACCENT, False, MONO)], 12, align=PP_ALIGN.CENTER, first=True)
    px = px + pw + Inches(0.18)
footer(s, 1)

# ============================================================
# SLIDE 2 — WELCOME
# ============================================================
s = slide(); chrome(s, 2)
kicker(s, "Welcome")
title_block(s, Inches(0.9), Inches(1.4), Inches(11.5), [
    [("Welcome to ", INK, True), ("TARS Lab Academy", ACCENT, True)]], size=38)
tb, tf = textbox(s, Inches(0.9), Inches(2.25), Inches(11), Inches(0.9))
para(tf, [("We don't teach you to watch. We teach you to ", MUTED, False),
          ("ship", ACCENT, True),
          (". Every day pairs a focused concept with a real build — by Day 45 you "
           "have a portfolio, not just notes.", MUTED, False)], 16, first=True, line=1.3)
cw, gap = Inches(3.7), Inches(0.3)
cx, cy, ch = Inches(0.9), Inches(3.5), Inches(2.6)
card(s, cx, cy, cw, ch, "🎯", "Outcome-first",
     "Each day ends with something running on your machine.")
card(s, cx+cw+gap, cy, cw, ch, "🧱", "Foundations → Frontier",
     "Python today; multi-agent systems & LLMOps later.")
card(s, cx+2*(cw+gap), cy, cw, ch, "🛠️", "Industry tools",
     "LangChain, LangGraph, RAG, AutoGen, MCP, Docker, K8s.")

# ============================================================
# SLIDE 3 — WHY NOW
# ============================================================
s = slide(); chrome(s, 3)
kicker(s, "Why this, why now")
title_block(s, Inches(0.9), Inches(1.4), Inches(11.5), [
    [("The most in-demand engineering", INK, True)],
    [("skill of the decade", INK, True)]], size=36)
tb, tf = textbox(s, Inches(0.9), Inches(3.0), Inches(11), Inches(1.0))
para(tf, [("Generative AI moved from research labs to every product team. Companies "
           "need engineers who take an LLM from a demo to a reliable, secure, "
           "monitored system.", MUTED, False)], 16, first=True, line=1.3)
stats = [("RAG", "Knowledge-grounded apps"), ("Agents", "Tools, memory, reasoning"),
         ("Eval", "Trust & quality gates"), ("Ops", "Deploy, scale, observe")]
sx = Inches(0.9)
for big, small in stats:
    tb, tf = textbox(s, sx, Inches(4.5), Inches(3.0), Inches(1.3))
    para(tf, [(big, ACCENT, True)], 40, first=True, space_after=2)
    para(tf, [(small, MUTED, False)], 13)
    sx = sx + Inches(3.05)

# ============================================================
# SLIDE 4 — WHAT YOU'LL BUILD
# ============================================================
s = slide(); chrome(s, 4)
kicker(s, "By Day 45 you can")
title_block(s, Inches(0.9), Inches(1.4), Inches(11.5),
            [[("What you'll be able to build", INK, True)]], size=38)
bullets(s, Inches(0.9), Inches(2.6), Inches(11.5), [
    [("Design and ship ", INK, False), ("hybrid RAG", ACCENT, True),
     (" systems with reranking and GraphRAG.", INK, False)],
    [("Build ", INK, False), ("autonomous agents", ACCENT, True),
     (" with LangGraph, CrewAI and AutoGen.", INK, False)],
    [("Fine-tune open models with ", INK, False), ("LoRA / QLoRA", ACCENT, True),
     (" on custom data.", INK, False)],
    [("Evaluate, guardrail and ", INK, False), ("secure", ACCENT, True),
     (" LLM apps for production.", INK, False)],
    [("Deploy with ", INK, False), ("FastAPI, Docker, Kubernetes", ACCENT, True),
     (" and full observability.", INK, False)],
], size=19, gap=14)

# ============================================================
# SLIDE 5 — THE MAP (12 phases)
# ============================================================
s = slide(); chrome(s, 5)
kicker(s, "The 45-day map")
title_block(s, Inches(0.9), Inches(1.4), Inches(11.5),
            [[("12 phases, one journey", INK, True)]], size=38)
phases = [
    ("P1", "Foundations (Python)", "D1–3"), ("P2", "How LLMs Work", "D4–7"),
    ("P3", "Building LLM Apps", "D8–11"), ("P4", "Advanced RAG", "D12–16"),
    ("P5", "AI Agents", "D17–21"), ("P6", "Multi-Agent (AutoGen)", "D22–25"),
    ("P7", "MCP & Coding Tools", "D26–29"), ("P8", "Fine-Tuning", "D30–32"),
    ("P9", "Evaluation", "D33–35"), ("P10", "Security & Guardrails", "D36–38"),
    ("P11", "LLMOps & Production", "D39–42"), ("P12", "Capstone & Career", "D43–45"),
]
colw, rowh = Inches(5.7), Inches(0.62)
x0, y0 = Inches(0.9), Inches(2.55)
for idx, (pn, pt, pd) in enumerate(phases):
    col, row = idx % 2, idx // 2
    x = x0 + col * (colw + Inches(0.4))
    y = y0 + row * rowh
    rect(s, x, y + rowh - Pt(2), colw, Pt(1), fill=LINE, radius=False)
    tb, tf = textbox(s, x, y, colw, rowh, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(f"{pn}   ", ACCENT, True, MONO), (pt, INK, True)], 14, first=True)
    tb2, tf2 = textbox(s, x, y, colw, rowh, anchor=MSO_ANCHOR.MIDDLE)
    para(tf2, [(pd, MUTED, False, MONO)], 12, align=PP_ALIGN.RIGHT, first=True)

# ============================================================
# SLIDE 6 — DAILY RHYTHM
# ============================================================
s = slide(); chrome(s, 6)
kicker(s, "The rhythm")
title_block(s, Inches(0.9), Inches(1.4), Inches(11.5),
            [[("How every day works", INK, True)]], size=38)
cw, gap, ch, cy = Inches(3.7), Inches(0.3), Inches(2.4), Inches(2.7)
card(s, Inches(0.9), cy, cw, ch, "📖", "1 · Learn",
     "A tight concept block — theory + intuition, no fluff.")
card(s, Inches(0.9)+cw+gap, cy, cw, ch, "⌨️", "2 · Code along",
     "A guided notebook you run line by line with us.")
card(s, Inches(0.9)+2*(cw+gap), cy, cw, ch, "🔨", "3 · Build",
     "A daily challenge you finish on your own.")
tb, tf = textbox(s, Inches(0.9), Inches(5.5), Inches(11.5), Inches(0.6))
para(tf, [("Today's notebook:  ", MUTED, False),
          ("Day01_Python_Core_and_Environments.ipynb", ACCENT, True, MONO)],
     16, first=True)

# ============================================================
# SLIDE 7 — DIVIDER: DAY 01
# ============================================================
s = slide(); chrome(s, 7)
kicker(s, "Phase 1 · Foundations")
tb, tf = textbox(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(2.4),
                 anchor=MSO_ANCHOR.MIDDLE)
para(tf, [("Day 01", ACCENT, True)], 110, first=True, space_after=0)
title_block(s, Inches(0.95), Inches(4.4), Inches(11.5),
            [[("Python Core & Environments", INK, True)]], size=36)
tb, tf = textbox(s, Inches(0.95), Inches(5.2), Inches(10.5), Inches(0.8))
para(tf, [("The language every GenAI tool speaks — and the workspace you'll live in "
           "for 45 days.", MUTED, False)], 16, first=True, line=1.3)

# ============================================================
# SLIDE 8 — DAY 1 AGENDA
# ============================================================
s = slide(); chrome(s, 8)
kicker(s, "Today's agenda")
title_block(s, Inches(0.9), Inches(1.4), Inches(11.5),
            [[("What we cover in Day 1", INK, True)]], size=38)
cw, gap, ch = Inches(5.7), Inches(0.4), Inches(1.85)
x0, y0 = Inches(0.9), Inches(2.7)
ag = [("⚙️", "Environments", "VS Code · Conda · UV · virtual environments."),
      ("🔤", "Language core", "Syntax, variables, datatypes, operators."),
      ("🔁", "Control flow", "Conditionals & loops."),
      ("📦", "Data structures", "Lists, comprehensions, tuples, sets, dicts.")]
for i, (ic, t, b) in enumerate(ag):
    col, row = i % 2, i // 2
    card(s, x0 + col*(cw+gap), y0 + row*(ch+Inches(0.3)), cw, ch, ic, t, b)

# ============================================================
# SLIDE 9 — WHY PYTHON
# ============================================================
s = slide(); chrome(s, 9)
kicker(s, "Foundation")
title_block(s, Inches(0.9), Inches(1.4), Inches(11.5),
            [[("Why Python for GenAI?", INK, True)]], size=38)
bullets(s, Inches(0.9), Inches(2.7), Inches(11.5), [
    [("It's the lingua franca", ACCENT, True),
     (" — PyTorch, Transformers, LangChain, every SDK is Python-first.", INK, False)],
    [("Readable & fast to prototype", ACCENT, True),
     (" — idea to running code in minutes.", INK, False)],
    [("Massive ecosystem", ACCENT, True),
     (" — data, ML, web, deployment, all in one language.", INK, False)],
    [("Glue for everything", ACCENT, True),
     (" — call APIs, orchestrate agents, serve models.", INK, False)],
], size=19, gap=16)

# ============================================================
# SLIDE 10 — ENVIRONMENTS
# ============================================================
s = slide(); chrome(s, 10)
kicker(s, "Setup")
title_block(s, Inches(0.9), Inches(1.4), Inches(11.5),
            [[("Your workspace: VS Code · Conda · UV", INK, True)]], size=32)
cw, gap, ch, cy = Inches(3.7), Inches(0.3), Inches(3.1), Inches(2.6)
card(s, Inches(0.9), cy, cw, ch, "🧩", "VS Code",
     "Editor + Python & Jupyter extensions. Where you write and run everything.")
card(s, Inches(0.9)+cw+gap, cy, cw, ch, "🐍", "Conda",
     "Manages Python versions + heavyweight ML packages cleanly.")
card(s, Inches(0.9)+2*(cw+gap), cy, cw, ch, "⚡", "UV",
     "Ultra-fast package & project manager — pip, but seconds not minutes.")

# ============================================================
# SLIDE 11 — VENVS + CODE
# ============================================================
s = slide(); chrome(s, 11)
kicker(s, "Isolation")
title_block(s, Inches(0.9), Inches(1.4), Inches(11.5),
            [[("Virtual environments — and why they matter", INK, True)]], size=30)
tb, tf = textbox(s, Inches(0.9), Inches(2.25), Inches(11.3), Inches(0.8))
para(tf, [("Every project gets its own sandbox of dependencies, so versions never "
           "collide. Non-negotiable in AI work.", MUTED, False)], 15, first=True,
     line=1.3)
codebox(s, Inches(0.9), Inches(3.3), Inches(11.5), Inches(3.0), [
    [("# Conda", MUTED)],
    [("conda create ", INK), ("-n ", INDIGO), ("genai ", AMBER),
     ("python=3.11 ", INK), ("-y", INDIGO)],
    [("conda activate ", INK), ("genai", AMBER)],
    [("", INK)],
    [("# UV (fast alternative)", MUTED)],
    [("uv venv ", INK), ("&& ", ACCENT), ("source ", ACCENT),
     (".venv/bin/activate", INK)],
    [("uv pip install ", INK), ("numpy pandas streamlit", AMBER)],
])

# ============================================================
# SLIDE 12 — VARIABLES / TYPES / OPERATORS
# ============================================================
s = slide(); chrome(s, 12)
kicker(s, "Language core")
title_block(s, Inches(0.9), Inches(1.4), Inches(11.5),
            [[("Variables, datatypes & operators", INK, True)]], size=34)
codebox(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(2.7), [
    [("name ", INK), ("= ", ACCENT), ('"TARS"', AMBER), ("          # str", MUTED)],
    [("days ", INK), ("= ", ACCENT), ("45", INDIGO), ("               # int", MUTED)],
    [("progress ", INK), ("= ", ACCENT), ("2.2", INDIGO),
     ("           # float", MUTED)],
    [("shipping ", INK), ("= ", ACCENT), ("True", INDIGO),
     ("          # bool", MUTED)],
    [("", INK)],
    [("# + - * / // % **   and/or/not   == != < >", MUTED)],
    [("done ", INK), ("= ", ACCENT), ("(days == ", INK), ("45", INDIGO),
     (") and shipping", INK)],
])
tb, tf = textbox(s, Inches(0.9), Inches(5.5), Inches(11.3), Inches(0.8))
para(tf, [("Dynamic typing, but types still matter — we add validation on Day 3 with "
           "Pydantic.", MUTED, False)], 15, first=True, line=1.3)

# ============================================================
# SLIDE 13 — CONTROL FLOW
# ============================================================
s = slide(); chrome(s, 13)
kicker(s, "Control flow")
title_block(s, Inches(0.9), Inches(1.4), Inches(11.5),
            [[("Conditionals & loops", INK, True)]], size=38)
codebox(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(3.5), [
    [("for ", ACCENT), ("day ", INK), ("in ", ACCENT), ("range(", INK),
     ("1", INDIGO), (", ", INK), ("46", INDIGO), ("):", INK)],
    [("    if ", ACCENT), ("day <= ", INK), ("3", INDIGO), (":", INK)],
    [("        phase ", INK), ("= ", ACCENT), ('"Foundations"', AMBER)],
    [("    elif ", ACCENT), ("day <= ", INK), ("7", INDIGO), (":", INK)],
    [("        phase ", INK), ("= ", ACCENT), ('"How LLMs Work"', AMBER)],
    [("    else", ACCENT), (":", INK)],
    [("        continue", ACCENT)],
])

# ============================================================
# SLIDE 14 — DATA STRUCTURES
# ============================================================
s = slide(); chrome(s, 14)
kicker(s, "Data structures")
title_block(s, Inches(0.9), Inches(1.4), Inches(11.5),
            [[("Lists · Tuples · Sets · Dicts", INK, True)]], size=36)
cw, gap, ch, cy = Inches(2.75), Inches(0.25), Inches(1.95), Inches(2.55)
ds = [("📋", "List", "Ordered, mutable", "[1, 2, 3]"),
      ("📌", "Tuple", "Ordered, fixed", "(x, y)"),
      ("🎲", "Set", "Unique, unordered", "{1, 2}"),
      ("🗂️", "Dict", "Key → value", '{"k": 1}')]
x = Inches(0.9)
for ic, t, b, code in ds:
    rect(s, x, cy, cw, ch, fill=BG2, line_color=LINE)
    tb, tf = textbox(s, x+Inches(0.2), cy+Inches(0.18), cw-Inches(0.4), ch-Inches(0.3))
    para(tf, [(ic, INK, False)], 20, first=True, space_after=3)
    para(tf, [(t, INK, True)], 15, space_after=2)
    para(tf, [(b, MUTED, False)], 11.5, space_after=6)
    para(tf, [(code, ACCENT, False, MONO)], 12.5)
    x = x + cw + gap
codebox(s, Inches(0.9), Inches(4.9), Inches(11.5), Inches(1.2), [
    [("squares ", INK), ("= ", ACCENT), ("[n", INK), ("**", ACCENT), ("2 ", INDIGO),
     ("for ", ACCENT), ("n ", INK), ("in ", ACCENT), ("range(", INK), ("10", INDIGO),
     (") ", INK), ("if ", ACCENT), ("n % ", INK), ("2 ", INDIGO), ("== ", ACCENT),
     ("0", INDIGO), ("]   # comprehension", MUTED)],
])

# ============================================================
# SLIDE 15 — BUILD
# ============================================================
s = slide(); chrome(s, 15)
kicker(s, "🔨 Today's build")
title_block(s, Inches(0.9), Inches(1.4), Inches(11.5),
            [[("Data-manipulation exercises", INK, True)]], size=36)
tb, tf = textbox(s, Inches(0.9), Inches(2.3), Inches(11.3), Inches(0.7))
para(tf, [("Open the notebook and work through it with us — then finish the challenge "
           "cells solo.", MUTED, False)], 16, first=True, line=1.3)
bullets(s, Inches(0.9), Inches(3.3), Inches(11.5), [
    [("Clean and transform a list of student records.", INK, False)],
    [("Use comprehensions to filter, map and aggregate.", INK, False)],
    [("De-duplicate with sets; index with dicts.", INK, False)],
    [("Stretch: ", AMBER, True),
     ("compute per-phase day counts from the syllabus.", INK, False)],
], size=18, gap=13)

# ============================================================
# SLIDE 16 — CLOSE
# ============================================================
s = slide(); chrome(s, 16)
kicker(s, "Day 1 complete")
tb, tf = textbox(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.6),
                 anchor=MSO_ANCHOR.MIDDLE)
para(tf, [("Let's build.", ACCENT, True)], 80, first=True)
tb, tf = textbox(s, Inches(0.9), Inches(4.4), Inches(11.0), Inches(1.4))
para(tf, [("Next:  ", MUTED, False), ("Day 2 — Python OOP & Advanced.", ACCENT, True)],
     18, first=True, line=1.35, space_after=8)
para(tf, [("Functions, classes, decorators, generators & a small OOP CLI project.",
           MUTED, False)], 16, line=1.3, space_after=10)
para(tf, [("Questions? Bring them to the notebook session.", MUTED, False)], 15)

# ---------- save ----------
out = "slides/TARS_Day01_Intro_Python_Core.pptx"
prs.save(out)
print(f"✅ saved {out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
